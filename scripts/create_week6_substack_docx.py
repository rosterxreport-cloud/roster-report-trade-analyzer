import csv
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Rebuild after all-team Week 6 passing-volume reconciliation.
src=Path("data/projections/2026_week6_league_v4_full.csv")
out=Path("data/projections/Roster_Report_Week6_Substack_Ready.docx")
with src.open(newline="",encoding="utf-8") as f: players=list(csv.DictReader(f))
doc=Document()
sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
normal=doc.styles["Normal"];normal.font.name="Arial";normal.font.size=Pt(9);normal.font.color.rgb=RGBColor(28,36,49)
normal.paragraph_format.space_after=Pt(4)
def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr();sh=OxmlElement("w:shd");sh.set(qn("w:fill"),fill);tcPr.append(sh)
def fmt(v,d=1):return f"{float(v):.{d}f}"
def line(p):
 if p["Position"]=="QB":
  return f'{fmt(p["Pass_Att"])} pass att · {fmt(p["Pass_Yds"],0)} pass yds · {fmt(p["Pass_TD"],2)} pass TD · {fmt(p["INT"],2)} INT · {fmt(p["Rush_Att"])} rush att · {fmt(p["Rush_Yds"],0)} rush yds · {fmt(p["Rush_TD"],2)} rush TD'
 return f'{fmt(p["Rush_Att"])} rush att · {fmt(p["Rush_Yds"],0)} rush yds · {fmt(p["Rush_TD"],2)} rush TD · {fmt(p["Targets"])} tgt · {fmt(p["Receptions"])} rec · {fmt(p["Rec_Yds"],0)} rec yds · {fmt(p["Rec_TD"],2)} rec TD'
doc.add_heading("Week 6 Fantasy Football Projections",0)
doc.add_paragraph(f'THE ROSTER REPORT | 2026 NFL Season | {len(players)} players | Standard, Half PPR and PPR')
doc.add_paragraph("Model estimates. Check final injury reports and active rosters before setting lineups.")
for pos,title in [("QB","Quarterbacks"),("RB","Running Backs"),("WR","Wide Receivers"),("TE","Tight Ends")]:
 arr=sorted((p for p in players if p["Position"]==pos),key=lambda p:-float(p["PPR"]))
 doc.add_heading(f"{title} ({len(arr)} players)",level=1)
 doc.add_paragraph("Ranked by projected PPR fantasy points.")
 t=doc.add_table(rows=1,cols=4);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=[Inches(4.9),Inches(.62),Inches(.62),Inches(.62)]
 for cell,w,label in zip(t.rows[0].cells,widths,["Player / Projected stat line","STD","Half","PPR"]):
  cell.width=w;shade(cell,"1F4178");cell.text=label
  for run in cell.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(9)
 for i,p in enumerate(arr):
  cells=t.add_row().cells
  for cell,w in zip(cells,widths):cell.width=w;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  if i%2: 
   for cell in cells:shade(cell,"F2F5F8")
  para=cells[0].paragraphs[0];para.paragraph_format.space_after=Pt(2)
  run=para.add_run(p["Player"]);run.bold=True;run.font.size=Pt(10)
  para=cells[0].add_paragraph();para.paragraph_format.space_after=Pt(0)
  run=para.add_run(line(p));run.font.color.rgb=RGBColor(101,117,137);run.font.size=Pt(8)
  for cell,key in zip(cells[1:],["Standard","Half_PPR","PPR"]):
   cell.text=fmt(p[key],2)
   cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
   for run in cell.paragraphs[0].runs:
    run.font.size=Pt(9)
    if key=="PPR":run.bold=True
doc.add_paragraph("Follow The Roster Report For More.")
out.parent.mkdir(parents=True,exist_ok=True)
doc.save(out)
print("CREATED",out,len(players))
