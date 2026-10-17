import csv
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src=Path('data/projections/2026_week6_2026_target_share_review.csv')
rows=list(csv.DictReader(src.open(encoding='utf-8-sig',newline='')))
D=Document(); sec=D.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
normal=D.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(9)
D.add_heading('Week 6 Fantasy Football Projections',0)
D.add_paragraph('THE ROSTER REPORT  |  2026 NFL SEASON  |  '+str(len(rows))+' PLAYERS  |  STD • HALF PPR • PPR')
D.add_paragraph('Model estimates. Check injury reports and active rosters before setting lineups.')
def fmt(v,d=1):return f'{float(v or 0):.{d}f}'
def stats(r):
 p=r['Position'];rush=f\"{fmt(r['Rush_Att'])} rush att · {fmt(r['Rush_Yds'],0)} rush yds · {fmt(r['Rush_TD'],2)} rush TD\";rec=f\"{fmt(r['Targets'])} tgt · {fmt(r['Receptions'])} rec · {fmt(r['Rec_Yds'],0)} rec yds · {fmt(r['Rec_TD'],2)} rec TD\"
 if p=='QB':return f\"{fmt(r['Pass_Att'])} pass att · {fmt(r['Pass_Yds'],0)} pass yds · {fmt(r['Pass_TD'],2)} pass TD · {fmt(r['INT'],2)} INT · \"+rush
 return rush+' · '+rec if p=='RB' else rec+(' · '+rush if float(r['Rush_Att'] or 0)>0 else '')
def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)
for pos,label in [('QB','Quarterbacks'),('RB','Running Backs'),('WR','Wide Receivers'),('TE','Tight Ends')]:
 group=sorted([r for r in rows if r['Position']==pos],key=lambda r:float(r['PPR'] or r['Standard'] or 0),reverse=True)
 D.add_heading(f'{label} ({len(group)} players)',level=2)
 t=D.add_table(rows=1,cols=4);t.autofit=False
 for c,w in zip(t.columns,[Inches(4.85),Inches(.67),Inches(.67),Inches(.67)]):c.width=w
 for c,txt in zip(t.rows[0].cells,['Player / Projected stat line','STD','Half','PPR']):
  c.text=txt;shade(c,'16355A');
  for p in c.paragraphs:
   for run in p.runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(9)
 for i,r in enumerate(group):
  cells=t.add_row().cells
  if i%2: [shade(c,'F0F3F7') for c in cells]
  p=cells[0].paragraphs[0];run=p.add_run(r['Player']);run.bold=True;run.font.size=Pt(9)
  p=cells[0].add_paragraph();p.paragraph_format.space_after=Pt(0);run=p.add_run(stats(r));run.font.size=Pt(7);run.font.color.rgb=RGBColor(90,104,122)
  for j,k in enumerate(['Standard','Half_PPR','PPR'],1):
   cells[j].text=fmt(r[k],2) if r[k] else '—'
   for p in cells[j].paragraphs:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
  for c in cells:
   tcPr=c._tc.get_or_add_tcPr();cant=OxmlElement('w:cantSplit');tcPr.append(cant)
D.save('data/projections/Roster_Report_Week6_Substack_Updated.docx')
print(len(rows))
